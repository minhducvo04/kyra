"""Approved workflows, deterministic documents and local review decisions."""
import json
import os
import re
import tempfile
from dataclasses import asdict, dataclass
from datetime import UTC, datetime
from pathlib import Path
from uuid import uuid4

from sqlalchemy import insert, select, update

from companion.db import engine_for_store
from companion.doc_qa import Finding, PandocRenderer, Renderer, inspect_docx
from companion.father_draft import DraftBrief, Drafter, DraftRejected, check_draft
from companion.paths import write_json
from companion.schema import father_tasks


@dataclass(frozen=True)
class WorkflowVersion:
    slug: str
    version: int
    title: str
    author: str
    template_path: Path
    required_facts: list[str]
    approved_at: str | None
    draft_instructions: str | None = None


class WorkflowStore:
    def __init__(self, root: Path):
        self.root = Path(root) / "workflows"

    def _directory(self, slug: str) -> Path:
        if not re.fullmatch(r"[a-z0-9]+(?:-[a-z0-9]+)*", slug):
            raise ValueError("Invalid workflow slug")
        return self.root / slug

    def save(self, version: WorkflowVersion) -> Path:
        if type(version.version) is not int or version.version < 1:
            raise ValueError("Workflow version must be a positive integer")
        directory = self._directory(version.slug)
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"v{version.version}.json"
        data = asdict(version)
        data["template_path"] = str(version.template_path.resolve(strict=True))
        # Publish a complete JSON file without ever replacing an existing version,
        # including when two writers race. write_json owns the temporary write.
        with tempfile.TemporaryDirectory(dir=directory) as scratch:
            staged = Path(scratch) / "version.json"
            write_json(staged, data)
            try:
                os.link(staged, path)
            except FileExistsError as exc:
                raise ValueError("Workflow version already exists") from exc
        return path

    def get(self, slug: str, version: int) -> WorkflowVersion | None:
        path = self._directory(slug) / f"v{version}.json"
        if not path.is_file():
            return None
        data = json.loads(path.read_text(encoding="utf-8"))
        data["template_path"] = Path(data["template_path"])
        result = WorkflowVersion(**data)
        if result.slug != slug or result.version != version:
            raise ValueError("Workflow identity does not match its file")
        return result

    def active(self, slug: str) -> WorkflowVersion | None:
        versions = []
        for path in self._directory(slug).glob("v*.json"):
            match = re.fullmatch(r"v([1-9][0-9]*)\.json", path.name)
            if match:
                version = self.get(slug, int(match[1]))
                if version and version.approved_at:
                    versions.append(version)
        return max(versions, key=lambda v: v.version, default=None)

    def list(self) -> list[WorkflowVersion]:
        return [version for directory in sorted(self.root.glob("*"))
                if directory.is_dir() and (version := self.active(directory.name))]


def build_document(version: WorkflowVersion, facts: dict[str, str], out_path: Path) -> Path:
    from docx import Document

    template = version.template_path.read_text(encoding="utf-8")
    text = re.sub(r"\{\{([^{}]+)\}\}", lambda match: facts[match[1].strip()], template)
    document = Document()
    document.core_properties.author = version.author
    document.core_properties.last_modified_by = version.author
    for line in text.splitlines():
        if line.strip():
            document.add_paragraph(line)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(out_path)
    return out_path


@dataclass(frozen=True)
class Task:
    id: int
    slug: str
    version: int
    status: str
    facts: dict[str, str]
    document_path: str
    report: dict
    note: str
    created_at: str
    decided_at: str | None


class ApprovalRefused(ValueError):
    """A document still has findings that prevent approval."""


class FatherTaskStore:
    def __init__(self, root: Path, workflows: WorkflowStore, *, renderer: Renderer | None = None,
                 out_dir: Path | None = None, drafter: Drafter | None = None):
        root = Path(root).resolve()
        # Explicit path prevents a personal DATABASE_URL from joining the tenants.
        self.engine = engine_for_store(root / "father.db", explicit=root / "father.db")
        self.workflows = workflows
        self.renderer = renderer if renderer is not None else PandocRenderer()
        self.drafter = drafter
        self.out_dir = Path(out_dir) if out_dir is not None else root / "father_tasks"

    @staticmethod
    def _task(row) -> Task:
        data = dict(row)
        data["facts"] = json.loads(data.pop("facts_json"))
        data["report"] = json.loads(data.pop("report_json"))
        return Task(**data)

    def get(self, task_id: int) -> Task | None:
        with self.engine.connect() as conn:
            row = conn.execute(select(father_tasks).where(father_tasks.c.id == task_id)).mappings().first()
        return self._task(row) if row else None

    def list(self) -> list[Task]:
        with self.engine.connect() as conn:
            rows = conn.execute(select(father_tasks).order_by(father_tasks.c.id.desc())).mappings()
            return [self._task(row) for row in rows]

    def _build(self, version: WorkflowVersion, facts: dict[str, str]) -> dict:
        template = version.template_path.read_text(encoding="utf-8")
        placeholders = {key.strip() for key in re.findall(r"\{\{([^{}]+)\}\}", template)}
        required = set(version.required_facts) | (placeholders - {"draft"})
        missing = sorted(key for key in required if not facts.get(key, "").strip())
        values = dict(facts)
        for key in missing:
            values[key] = ""
        draft_provider = None
        if "draft" in placeholders:
            if self.drafter is None:
                raise ValueError("drafter_not_configured")
            draft = self.drafter.draft(DraftBrief(version.title, dict(facts), version.draft_instructions or ""))
            problems = check_draft(draft, values)
            if problems:
                raise DraftRejected(problems)
            values["draft"] = draft
            draft_provider = self.drafter.provider
        document = build_document(version, values, self.out_dir / uuid4().hex / "document.docx")
        report = inspect_docx(document, facts=facts, allowed_authors={version.author}, renderer=self.renderer)
        report.findings.extend(Finding("fact_missing", "facts", key) for key in missing)
        report_data = asdict(report)
        report_data["rendered_pages"] = [str(path) for path in report.rendered_pages]
        report_data["draft_provider"] = draft_provider
        return {"document_path": str(document), "report_json": json.dumps(report_data),
                "facts_json": json.dumps(facts)}

    def start_task(self, slug: str, facts: dict[str, str]) -> Task:
        version = self.workflows.active(slug)
        if version is None:
            raise ValueError("No approved workflow version")
        values = self._build(version, facts)
        with self.engine.begin() as conn:
            result = conn.execute(insert(father_tasks).values(
                **values, slug=slug, version=version.version, status="review", note="",
                created_at=datetime.now(UTC).isoformat(), decided_at=None,
            ))
            task_id = result.inserted_primary_key[0]
        return self.get(task_id)

    def update_facts(self, task_id: int, facts: dict[str, str]) -> Task:
        task = self.get(task_id)
        if task is None:
            raise KeyError(task_id)
        if task.status != "review":
            raise ValueError("Only a task awaiting review can be edited; start a new task")
        version = self.workflows.get(task.slug, task.version)
        if version is None or not version.approved_at:
            raise ValueError("Approved workflow version is missing")
        values = self._build(version, facts)
        with self.engine.begin() as conn:
            result = conn.execute(update(father_tasks).where(
                father_tasks.c.id == task_id, father_tasks.c.status == "review",
                father_tasks.c.report_json == json.dumps(task.report),
            ).values(**values))
            if result.rowcount != 1:
                raise ValueError("Task changed during review; reload it")
        return self.get(task_id)

    def decide(self, task_id: int, decision: str, note: str = "") -> Task:
        if decision not in {"approve", "changes_requested"}:
            raise ValueError("Unknown decision")
        task = self.get(task_id)
        if task is None:
            raise KeyError(task_id)
        if task.status != "review":
            raise ValueError("This task already has a decision")
        if decision == "approve":
            blocked = {finding["kind"] for finding in task.report["findings"]} & {
                "dash", "provider_name", "fact_missing", "truncation",
            }
            if task.report["render_error"] or not task.report["rendered_pages"]:
                blocked.add("render_error")
            if blocked:
                raise ApprovalRefused("Approval refused: " + ", ".join(sorted(blocked)))
        with self.engine.begin() as conn:
            result = conn.execute(update(father_tasks).where(
                father_tasks.c.id == task_id, father_tasks.c.status == "review",
                father_tasks.c.report_json == json.dumps(task.report),
            ).values(status="approved" if decision == "approve" else decision,
                     note=note, decided_at=datetime.now(UTC).isoformat()))
            if result.rowcount != 1:
                raise ValueError("Task changed during review; reload it")
        return self.get(task_id)
