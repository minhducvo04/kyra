"""Father's task cards on synthetic data (red until Codex builds F03).

Plan: docs/plans/2026-09-16-father.md, section F03. No model call anywhere in this slice: an approved
workflow version plus a facts table goes through a deterministic python-docx builder, then doc_qa.

CONTRACT (companion/father.py)
  @dataclass(frozen=True) WorkflowVersion: slug, version: int, title, author, template_path: Path,
      required_facts: list[str], approved_at: str | None
  class WorkflowStore:
      __init__(root: Path)                        # <root>/workflows/<slug>/v<N>.json, one file per version
      save(version) -> Path                        # refuses to overwrite an existing file (ValueError)
      active(slug) -> WorkflowVersion | None       # highest version with approved_at set
  build_document(version, facts: dict[str, str], out_path: Path) -> Path
      template is UTF-8 text with {{key}} placeholders; one paragraph per non-empty line;
      core author and last_modified_by = version.author; missing key -> KeyError
  @dataclass(frozen=True) Task: id, slug, version, status, facts, document_path, report: dict, note, created_at, decided_at
      statuses: "review" | "approved" | "changes_requested"
  class FatherTaskStore:
      __init__(root: Path, workflows: WorkflowStore, *, renderer, out_dir: Path | None = None)   # <root>/father.db
      start_task(slug, facts) -> Task              # builds, runs inspect_docx(allowed_authors={author}), status review
      get(task_id) -> Task | None;  list() -> list[Task] newest first
      decide(task_id, decision, note="") -> Task
          "approve" -> ApprovalRefused when report has a dash, provider_name, fact_missing or truncation
                       finding, or a render_error; the message names the kinds
          "changes_requested" -> stores the note; either decision sets decided_at
  class ApprovalRefused(ValueError)
  webapp: GET /father and GET /api/father/tasks -> 404 "not_found" unless get_settings().tenant == "father"
"""
import json
from pathlib import Path

import pytest

from companion.settings import get_settings


@pytest.fixture
def fa():
    import companion.father as father

    return father


class OnePage:
    def render(self, path, out_dir):
        page = Path(out_dir) / "page-1.png"
        page.parent.mkdir(parents=True, exist_ok=True)
        page.write_bytes(b"\x89PNG fake")
        return [page]


TEMPLATE = "Dear {{recipient}},\n\nThe Northwind invoice for {{month}} totals {{amount}}.\n\nKind regards,\nSam\n"


def _workflow(fa, root, *, version=1, approved="2026-09-16T00:00:00+00:00", template=TEMPLATE):
    tpl = root / f"template-v{version}.txt"
    tpl.write_text(template, encoding="utf-8")
    v = fa.WorkflowVersion(slug="monthly-invoice", version=version, title="Monthly invoice letter", author="Northwind",
                           template_path=tpl, required_facts=["recipient", "month", "amount"], approved_at=approved)
    fa.WorkflowStore(root).save(v)
    return v


FACTS = {"recipient": "Alex Rivera", "month": "March", "amount": "1,250.00"}


def test_active_version_is_the_highest_approved_and_files_are_never_overwritten(fa, tmp_path):
    _workflow(fa, tmp_path, version=1)
    _workflow(fa, tmp_path, version=2, approved=None)  # drafted, not approved
    store = fa.WorkflowStore(tmp_path)
    assert store.active("monthly-invoice").version == 1
    assert json.loads((tmp_path / "workflows" / "monthly-invoice" / "v1.json").read_text())["author"] == "Northwind"
    with pytest.raises(ValueError):
        store.save(_workflow(fa, tmp_path / "again", version=1))  # same slug and version into the first store
    assert store.active("nope") is None


def test_builder_fills_placeholders_and_stamps_the_author(fa, tmp_path):
    from docx import Document

    v = _workflow(fa, tmp_path)
    out = fa.build_document(v, FACTS, tmp_path / "letter.docx")
    d = Document(out)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Alex Rivera" in text and "1,250.00" in text and "{{" not in text
    assert d.core_properties.author == "Northwind" and d.core_properties.last_modified_by == "Northwind"
    with pytest.raises(KeyError):
        fa.build_document(v, {"recipient": "Alex Rivera"}, tmp_path / "short.docx")


def test_start_task_builds_checks_and_waits_for_review(fa, tmp_path):
    _workflow(fa, tmp_path)
    tasks = fa.FatherTaskStore(tmp_path, fa.WorkflowStore(tmp_path), renderer=OnePage())
    task = tasks.start_task("monthly-invoice", FACTS)
    assert task.status == "review" and task.version == 1 and Path(task.document_path).exists()
    assert task.report["findings"] == [] and task.report["render_error"] is None and len(task.report["rendered_pages"]) == 1
    assert tasks.list()[0].id == task.id and tasks.get(task.id).facts == FACTS


def test_approval_is_refused_while_the_report_has_hard_findings(fa, tmp_path):
    _workflow(fa, tmp_path, template="Dear {{recipient}} — see the {{month}} invoice for {{amount}}. Drafted by Claude.\n")
    tasks = fa.FatherTaskStore(tmp_path, fa.WorkflowStore(tmp_path), renderer=OnePage())
    task = tasks.start_task("monthly-invoice", FACTS)
    kinds = {f["kind"] for f in task.report["findings"]}
    assert {"dash", "provider_name"} <= kinds
    with pytest.raises(fa.ApprovalRefused, match="dash"):
        tasks.decide(task.id, "approve")
    changed = tasks.decide(task.id, "changes_requested", note="remove the dash and the attribution")
    assert changed.status == "changes_requested" and changed.note.startswith("remove") and changed.decided_at


def test_clean_task_can_be_approved_and_nothing_is_sent(fa, tmp_path):
    _workflow(fa, tmp_path)
    tasks = fa.FatherTaskStore(tmp_path, fa.WorkflowStore(tmp_path), renderer=OnePage())
    task = tasks.decide(tasks.start_task("monthly-invoice", FACTS).id, "approve")
    assert task.status == "approved" and task.decided_at
    assert not hasattr(tasks, "send") and not hasattr(tasks, "deliver")


def test_father_routes_exist_only_for_the_father_tenant(monkeypatch, tmp_path):
    from fastapi.testclient import TestClient

    from companion import webapp

    monkeypatch.setenv("KYRA_API_TOKEN", "")
    monkeypatch.setenv("KYRA_TRUST_LOOPBACK", "true")
    monkeypatch.delenv("KYRA_TENANT", raising=False)
    get_settings.cache_clear()
    with TestClient(webapp.app, base_url="http://127.0.0.1:8420", client=("127.0.0.1", 4321)) as client:
        assert client.get("/api/father/tasks").status_code == 404
        assert client.get("/father").status_code == 404
        monkeypatch.setenv("KYRA_TENANT", "father")
        monkeypatch.setenv("KYRA_DATA_DIR", str(tmp_path / "father-data"))
        get_settings.cache_clear()
        assert client.get("/api/father/tasks").json() == {"tasks": []}
        assert "Start a task" in client.get("/father").text
    get_settings.cache_clear()
